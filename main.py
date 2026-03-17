from openpyxl import load_workbook
import warnings

from datetime import datetime
import os

# Usuwa niepotrzebne blendy
warnings.filterwarnings("ignore", category=UserWarning)


'''
Import Wykaz.xlsx
'''
Wykaz = "Wykaz.xlsx"
Wykaz_SHEET = "Szablon 23.11.0.0"
FINAL_DIR = "Final"

# Otwiera glowny plik wykazu i wybiera arkusz roboczy
main_wb = load_workbook(Wykaz)
main_ws = main_wb[Wykaz_SHEET]

# Dane w wykazie zaczynaja sie od wiersza 3, bo wyzej sa naglowki
WYKAZ_START_ROW = 3

# Mapowanie pol na kolumny w pliku Wykaz.xlsx
WYKAZ_KOLUMNY = {
    "nr_rej": 1,
    "rodzaj": 2,
    "marka": 3,
    "model": 4,
    "vin": 5,
    "rok_prod": 6,
    "pojemnosc": 7,
    "wartosc_pojazdu": 8,
    "poczatek_polisy": 19,
    "koniec_polisy": 20,
}

# Mapowanie pol na kolumny w plikach eksportow
EXPORT_KOLUMNY = {
    "wartosc_pojazdu": 1,
    "nr_rej": 2,
    "marka": 4,
    "typ": 5,
    "model": 6,
    "rodzaj": 7,
    "rok_prod": 9,
    "vin": 11,
    "pojemnosc": 12,
    "poczatek_polisy": 22,
    "koniec_polisy": 23,
}

'''
Testwoe wyswietlanie main
'''


def Test_main():
    for row in main_ws.iter_rows(min_row=1, max_row=10, min_col=1, max_col=10):
        for cell in row:
            print(cell.value, end=" ")
        print()

'''
Skanowanie i Wczytywanie Eksporty
'''
EXPORTS_DIR = "Eskporty"


def wczytaj_eksporty(folder=EXPORTS_DIR):
    eksporty = []

    if not os.path.isdir(folder):
        print(f"Folder z eksportami nie istnieje: {folder}")
        return eksporty

# Wczytuje wszystkie pliki .xlsx z folderu eksportow
    for nazwa_pliku in sorted(os.listdir(folder)):
        if nazwa_pliku.startswith("~$"):
            continue
        if not nazwa_pliku.lower().endswith(".xlsx"):
            continue

        sciezka = os.path.join(folder, nazwa_pliku)
        wb = load_workbook(sciezka, data_only=True)
        ws = wb.active

        eksporty.append({
            "nazwa": nazwa_pliku,
            "sciezka": sciezka,
            "workbook": wb,
            "worksheet": ws,
        })

    return eksporty

'''
NR KOlumny kota bedzie odczytywana
2 = kolumna B, 3 = kolumna C, 4 = kolumna D
'''
NR_KOLUMNY = 2

def drukuj_druga_kolumne(eksporty):
    for eksport in eksporty:
        ws = eksport["worksheet"]
        print(f"\nPlik: {eksport['nazwa']}")

        for nr_wiersza, row in enumerate(
            ws.iter_rows(min_col=NR_KOLUMNY, max_col=NR_KOLUMNY, values_only=True),
            start=1,
        ):
            wartosc = row[0]

            if wartosc is None:
                continue
            if str(wartosc).strip() == "":
                continue
            if str(wartosc).strip().lower() == "nr rej":
                continue

            print(f"Wiersz {nr_wiersza}: {wartosc}")


def normalizuj_tekst(wartosc):
# Ujednolica tekst do porownan, np. przy sprawdzaniu duplikatow
    if wartosc is None:
        return ""
    return str(wartosc).strip().lower()


def pierwsza_niepusta_wartosc(*wartosci):
    # Zwraca pierwsza niepusta wartosc z podanych pol
    for wartosc in wartosci:
        if wartosc is None:
            continue
        if str(wartosc).strip() == "":
            continue
        return wartosc
    return None


def pobierz_pojazdy_z_eksportow(eksporty):
    pojazdy = []

    for eksport in eksporty:
        ws = eksport["worksheet"]

        # Pomija pierwsze dwa wiersze, bo zawieraja tytul i naglowki eksportu
        for nr_wiersza, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            nr_rej = pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["nr_rej"] - 1])
            vin = pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["vin"] - 1])

            # Jesli nie ma ani numeru rejestracyjnego, ani VIN-u, to rekord jest pusty
            if nr_rej is None and vin is None:
                continue

            # Buduje ujednolicony slownik danych pojazdu z eksportu
            pojazdy.append({
                "nr_wiersza_eksport": nr_wiersza,
                "plik": eksport["nazwa"],
                "nr_rej": str(nr_rej).strip() if nr_rej is not None else None,
                "rodzaj": pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["rodzaj"] - 1]),
                "marka": pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["marka"] - 1]),
                "model": pierwsza_niepusta_wartosc(
                    row[EXPORT_KOLUMNY["model"] - 1],
                    row[EXPORT_KOLUMNY["typ"] - 1],
                ),
                "vin": str(vin).strip() if vin is not None else None,
                "rok_prod": pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["rok_prod"] - 1]),
                "pojemnosc": pierwsza_niepusta_wartosc(row[EXPORT_KOLUMNY["pojemnosc"] - 1]),
                "wartosc_pojazdu": pierwsza_niepusta_wartosc(
                    row[EXPORT_KOLUMNY["wartosc_pojazdu"] - 1]
                ),
                "poczatek_polisy": pierwsza_niepusta_wartosc(
                    row[EXPORT_KOLUMNY["poczatek_polisy"] - 1]
                ),
                "koniec_polisy": pierwsza_niepusta_wartosc(
                    row[EXPORT_KOLUMNY["koniec_polisy"] - 1]
                ),
            })

    return pojazdy


def pobierz_istniejace_klucze_z_wykazu():
    numery_rej = set()
    viny = set()

    # Zbiera istniejace numery rejestracyjne i VIN-y z wykazu do szybkiego porownania
    for row in main_ws.iter_rows(min_row=WYKAZ_START_ROW, values_only=True):
        nr_rej = normalizuj_tekst(row[WYKAZ_KOLUMNY["nr_rej"] - 1])
        vin = normalizuj_tekst(row[WYKAZ_KOLUMNY["vin"] - 1])

        if nr_rej:
            numery_rej.add(nr_rej)
        if vin:
            viny.add(vin)

    return numery_rej, viny


def znajdz_pierwszy_pusty_wiersz(ws, start_row, kolumna_sprawdzana):
    # Szuka pierwszego wolnego wiersza, gdzie mozna dopisac nowy pojazd
    for nr_wiersza in range(start_row, ws.max_row + 2):
        wartosc = ws.cell(row=nr_wiersza, column=kolumna_sprawdzana).value
        if wartosc in (None, ""):
            return nr_wiersza
    return ws.max_row + 1


def dodaj_pojazdy_do_wykazu(eksporty):
    # Pobiera dane z eksportow i aktualny stan wykazu
    pojazdy = pobierz_pojazdy_z_eksportow(eksporty)
    istniejące_numery_rej, istniejące_viny = pobierz_istniejace_klucze_z_wykazu()
    pierwszy_pusty_wiersz = znajdz_pierwszy_pusty_wiersz(
        main_ws,
        WYKAZ_START_ROW,
        WYKAZ_KOLUMNY["nr_rej"],
    )

    dodane = []
    pominięte = []

    for pojazd in pojazdy:
        klucz_nr_rej = normalizuj_tekst(pojazd["nr_rej"])
        klucz_vin = normalizuj_tekst(pojazd["vin"])

        # Rekord jest duplikatem, jesli numer rejestracyjny albo VIN juz istnieje w wykazie
        duplikat_nr_rej = klucz_nr_rej and klucz_nr_rej in istniejące_numery_rej
        duplikat_vin = klucz_vin and klucz_vin in istniejące_viny

        if duplikat_nr_rej or duplikat_vin:
            powod = []
            if duplikat_nr_rej:
                powod.append("nr rejestracyjny")
            if duplikat_vin:
                powod.append("VIN")

            pominięte.append(
                f"{pojazd['plik']} wiersz {pojazd['nr_wiersza_eksport']}: "
                f"pominięto, istnieje już {' i '.join(powod)}"
            )
            continue

        # Zapisuje nowy pojazd do odpowiednich kolumn w wykazie
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["nr_rej"], pojazd["nr_rej"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["rodzaj"], pojazd["rodzaj"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["marka"], pojazd["marka"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["model"], pojazd["model"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["vin"], pojazd["vin"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["rok_prod"], pojazd["rok_prod"])
        main_ws.cell(pierwszy_pusty_wiersz, WYKAZ_KOLUMNY["pojemnosc"], pojazd["pojemnosc"])
        main_ws.cell(
            pierwszy_pusty_wiersz,
            WYKAZ_KOLUMNY["wartosc_pojazdu"],
            pojazd["wartosc_pojazdu"],
        )
        main_ws.cell(
            pierwszy_pusty_wiersz,
            WYKAZ_KOLUMNY["poczatek_polisy"],
            pojazd["poczatek_polisy"],
        )
        main_ws.cell(
            pierwszy_pusty_wiersz,
            WYKAZ_KOLUMNY["koniec_polisy"],
            pojazd["koniec_polisy"],
        )
        '''
        Po dodaniu rekordu od razu dopisuje jego klucze do zbiorow
        zeby nie dodac duplikatu z kolejnego pliku eksportu
        '''
        if klucz_nr_rej:
            istniejące_numery_rej.add(klucz_nr_rej)
        if klucz_vin:
            istniejące_viny.add(klucz_vin)

        dodane.append((pierwszy_pusty_wiersz, pojazd["nr_rej"], pojazd["vin"]))
        pierwszy_pusty_wiersz += 1

    final_path = przygotuj_sciezke_wynikowa()
    main_wb.save(final_path)

    print(f"Dodano {len(dodane)} nowych pojazdow do wykazu.")
    print(f"Zapisano nowy plik wynikowy: {final_path}")
    for nr_wiersza, nr_rej, vin in dodane:
        print(f"Wiersz {nr_wiersza}: nr rej={nr_rej}, VIN={vin}")

    print(f"Pominieto {len(pominięte)} duplikatow.")
    for komunikat in pominięte:
        print(komunikat)


def przygotuj_sciezke_wynikowa():
    # Tworzy folder Final i zwraca sciezke dla nowego pliku wynikowego
    os.makedirs(FINAL_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    return os.path.join(FINAL_DIR, f"Wykaz_scalony_{timestamp}.xlsx")



if __name__ == '__main__':
    # Wczytuje eksporty i zapisuje scalony wynik jako nowy plik w folderze Final
    eksporty = wczytaj_eksporty()
    print(f"Scalanie eksportow do pliku: {Wykaz}")
    dodaj_pojazdy_do_wykazu(eksporty)
